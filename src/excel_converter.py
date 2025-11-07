"""
Excel Converter for Etsy Taxonomy Data

Converts JSON taxonomy data into formatted Excel spreadsheet.
"""

import json
from pathlib import Path
from typing import Dict, List, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelConverter:
    """Converts taxonomy JSON data to Excel format."""

    # Color scheme
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    REQUIRED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    ENUMERATED_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # Border style
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    def __init__(self):
        """Initialize the converter."""
        self.workbook = None
        self.properties_sheet = None
        self.enums_sheet = None
        self.summary_sheet = None

    def create_workbook(self) -> Workbook:
        """Create a new Excel workbook with sheets."""
        self.workbook = Workbook()

        # Remove default sheet
        if "Sheet" in self.workbook.sheetnames:
            del self.workbook["Sheet"]

        # Create sheets
        self.summary_sheet = self.workbook.create_sheet("Summary", 0)
        self.properties_sheet = self.workbook.create_sheet("Properties", 1)
        self.enums_sheet = self.workbook.create_sheet("Enumerated Values", 2)

        return self.workbook

    def write_summary_sheet(self, data: Dict[str, Any]) -> None:
        """Write summary information to the Summary sheet."""
        sheet = self.summary_sheet
        metadata = data["metadata"]

        # Title
        sheet["A1"] = "Etsy Taxonomy Metadata Summary"
        sheet["A1"].font = Font(bold=True, size=14)

        # Metadata
        row = 3
        info = [
            ("Taxonomy ID", metadata.get("taxonomy_id", "N/A")),
            ("Fetched At", metadata.get("fetched_at", "N/A")),
            ("Total Properties", metadata.get("total_properties", 0)),
            ("Required Properties", metadata.get("required_properties", 0)),
            ("Optional Properties", metadata.get("optional_properties", 0)),
            ("Enumerated Properties", metadata.get("enumerated_properties", 0)),
        ]

        for label, value in info:
            sheet[f"A{row}"] = label
            sheet[f"A{row}"].font = Font(bold=True)
            sheet[f"B{row}"] = value
            row += 1

        # Legend
        row += 2
        sheet[f"A{row}"] = "Legend:"
        sheet[f"A{row}"].font = Font(bold=True)

        row += 1
        sheet[f"A{row}"] = "Yellow highlight"
        sheet[f"A{row}"].fill = self.REQUIRED_FILL
        sheet[f"B{row}"] = "Required property"

        row += 1
        sheet[f"A{row}"] = "Green highlight"
        sheet[f"A{row}"].fill = self.ENUMERATED_FILL
        sheet[f"B{row}"] = "Has enumerated values"

        # Adjust column widths
        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 40

    def write_properties_sheet(self, properties: List[Dict[str, Any]]) -> None:
        """Write property data to the Properties sheet."""
        sheet = self.properties_sheet

        # Define headers
        headers = [
            "Property ID",
            "Property Name",
            "Display Name",
            "Required?",
            "Enumerated?",
            "Enum Count",
            "Description",
            "Supports Attributes",
            "Supports Variations",
            "Multi-valued?",
            "Max Values",
            "Has Scales?"
        ]

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.THIN_BORDER

        # Write data rows
        for row_idx, prop in enumerate(properties, start=2):
            # Determine row highlighting
            fill = None
            if prop.get("is_required"):
                fill = self.REQUIRED_FILL
            elif prop.get("is_enumerated"):
                fill = self.ENUMERATED_FILL

            # Write each column
            row_data = [
                prop.get("property_id", ""),
                prop.get("property_name", ""),
                prop.get("display_name", ""),
                "YES" if prop.get("is_required") else "NO",
                "YES" if prop.get("is_enumerated") else "NO",
                prop.get("enumerated_values_count", 0),
                prop.get("description", ""),
                "YES" if prop.get("supports_attributes") else "NO",
                "YES" if prop.get("supports_variations") else "NO",
                "YES" if prop.get("is_multivalued") else "NO",
                prop.get("max_values_allowed", ""),
                "YES" if prop.get("scales") else "NO"
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill

        # Adjust column widths
        column_widths = {
            "A": 12,  # Property ID
            "B": 25,  # Property Name
            "C": 25,  # Display Name
            "D": 10,  # Required
            "E": 12,  # Enumerated
            "F": 10,  # Enum Count
            "G": 40,  # Description
            "H": 15,  # Supports Attributes
            "I": 15,  # Supports Variations
            "J": 12,  # Multi-valued
            "K": 11,  # Max Values
            "L": 11   # Has Scales
        }

        for col_letter, width in column_widths.items():
            sheet.column_dimensions[col_letter].width = width

        # Freeze header row
        sheet.freeze_panes = "A2"

    def write_enumerated_values_sheet(self, properties: List[Dict[str, Any]]) -> None:
        """Write enumerated values to the Enumerated Values sheet."""
        sheet = self.enums_sheet

        # Headers
        headers = ["Property ID", "Property Name", "Value ID", "Value Name"]

        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.THIN_BORDER

        # Write enumerated values
        row_idx = 2
        for prop in properties:
            if not prop.get("is_enumerated"):
                continue

            property_id = prop.get("property_id", "")
            property_name = prop.get("property_name", "")
            enum_values = prop.get("enumerated_values", [])

            if not enum_values:
                continue

            for enum_val in enum_values:
                row_data = [
                    property_id,
                    property_name,
                    enum_val.get("value_id", ""),
                    enum_val.get("name", "")
                ]

                for col_idx, value in enumerate(row_data, start=1):
                    cell = sheet.cell(row=row_idx, column=col_idx)
                    cell.value = value
                    cell.border = self.THIN_BORDER
                    cell.alignment = Alignment(vertical="top")

                row_idx += 1

        # Adjust column widths
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 12
        sheet.column_dimensions["D"].width = 40

        # Freeze header row
        sheet.freeze_panes = "A2"

    def convert_json_to_excel(self, json_path: str, excel_path: str) -> None:
        """
        Convert JSON taxonomy data to Excel format.

        Args:
            json_path: Path to input JSON file
            excel_path: Path to output Excel file
        """
        # Load JSON data
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        # Create workbook and sheets
        self.create_workbook()

        # Write data to sheets
        self.write_summary_sheet(data)
        self.write_properties_sheet(data["properties"])
        self.write_enumerated_values_sheet(data["properties"])

        # Save workbook
        self.workbook.save(excel_path)
        print(f"✓ Excel saved to: {excel_path}")


def convert_to_excel(json_path: str, excel_path: str = None) -> str:
    """
    Convenience function to convert JSON to Excel.

    Args:
        json_path: Path to JSON file
        excel_path: Optional path to Excel file (auto-generated if not provided)

    Returns:
        Path to the created Excel file
    """
    if excel_path is None:
        # Auto-generate Excel filename
        json_file = Path(json_path)
        excel_path = json_file.with_suffix('.xlsx')

    converter = ExcelConverter()
    converter.convert_json_to_excel(json_path, excel_path)

    return excel_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python excel_converter.py <json_file> [output_excel_file]")
        sys.exit(1)

    json_file = sys.argv[1]
    excel_file = sys.argv[2] if len(sys.argv) > 2 else None

    convert_to_excel(json_file, excel_file)
